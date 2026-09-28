from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app import ImportDatabase, create_preview_workbook, inspect_workbook


def sample_workbook() -> bytes:
    workbook = Workbook()
    pivot = workbook.active
    pivot.title = "PIVOT"
    pivot["B5"] = "Тестова мережа"
    pivot["F1"] = "за тестовий період"

    for sheet_name in ("LA_Price", "Реквізити", "Реквізити_лік"):
        sheet = workbook.create_sheet(sheet_name)
        sheet["B1"] = "Клієнт"
        sheet["B2"] = "Тестова мережа"
    for sheet_name in ("акт_ЛІК", "акт_ДД"):
        sheet = workbook.create_sheet(sheet_name)
        sheet["C12"] = 100
        sheet["C13"] = 20
        sheet["C14"] = 120
    for sheet_name, total_cell in (
        ("звіт викладка+MHL(ЛІК)", "S2029"),
        ("звіт викладка+MHL(ДД)", "P2029"),
    ):
        sheet = workbook.create_sheet(sheet_name)
        sheet["A9"] = "Номер"
        sheet["B9"] = "Аптека"
        sheet[total_cell] = 100

    content = BytesIO()
    workbook.save(content)
    workbook.close()
    return content.getvalue()


class WorkbookTests(unittest.TestCase):
    def test_inspection_finds_current_network_and_period(self) -> None:
        info = inspect_workbook(sample_workbook())
        self.assertEqual(info.network, "Тестова мережа")
        self.assertEqual(info.period, "за тестовий період")
        self.assertEqual(info.networks, ["Тестова мережа"])

    def test_both_variants_create_two_sheet_workbooks(self) -> None:
        content = sample_workbook()
        for variant in ("ЛІК", "ДД"):
            with self.subTest(variant=variant):
                workbook, warnings = create_preview_workbook(
                    content, variant, "Тестова мережа", "за тестовий період"
                )
                self.assertEqual(workbook.sheetnames, ["Акт", "Звіт з викладки"])
                self.assertEqual(workbook["Акт"]["C14"].value, 120)
                self.assertEqual(warnings, [])
                workbook.close()

    def test_refuses_network_not_selected_in_pivot(self) -> None:
        with self.assertRaisesRegex(ValueError, "Дані PIVOT підготовлені"):
            create_preview_workbook(
                sample_workbook(), "ДД", "Інша мережа", "будь-який період"
            )

    def test_database_deduplicates_imports_and_persists_workbook(self) -> None:
        content = sample_workbook()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "imports.sqlite3"
            database = ImportDatabase(path)
            source = Path(directory) / "sample.xlsx"
            source.write_bytes(content)
            first_id, inserted = database.add_workbook(
                source, content, "Тестова мережа", "за тестовий період"
            )
            duplicate_id, inserted_again = database.add_workbook(
                source, content, "Тестова мережа", "за тестовий період"
            )
            self.assertTrue(inserted)
            self.assertFalse(inserted_again)
            self.assertEqual(first_id, duplicate_id)
            restored = database.get_workbook(first_id)
            workbook = load_workbook(BytesIO(restored), read_only=True, data_only=True)
            self.assertEqual(workbook["PIVOT"]["B5"].value, "Тестова мережа")
            workbook.close()


if __name__ == "__main__":
    unittest.main()
