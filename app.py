from __future__ import annotations

import hashlib
import os
import sqlite3
import tkinter as tk
import warnings
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


APP_NAME = "Генератор актів і звітів"
REQUIRED_SHEETS = {
    "PIVOT",
    "LA_Price",
    "Реквізити",
    "Реквізити_лік",
}
VARIANTS = {
    "ЛІК": {
        "requisites": "Реквізити_лік",
        "act": "акт_ЛІК",
        "report": "звіт викладка+MHL(ЛІК)",
        "report_total": "S2029",
    },
    "ДД": {
        "requisites": "Реквізити",
        "act": "акт_ДД",
        "report": "звіт викладка+MHL(ДД)",
        "report_total": "P2029",
    },
}
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
WHITE_BOLD = Font(color="FFFFFF", bold=True)
THIN_GREY = Side(style="thin", color="D9E1F2")


def _load_workbook_values(content: bytes):
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message=r"Print area cannot be set to Defined name: .*!\$A:\$G\.",
            category=UserWarning,
            module=r"openpyxl\.reader\.workbook",
        )
        return load_workbook(BytesIO(content), data_only=True, read_only=True)


def database_path() -> Path:
    if os.name == "nt":
        root = Path(os.environ.get("LOCALAPPDATA", Path.home() / "AppData/Local"))
    else:
        root = Path(os.environ.get("XDG_DATA_HOME", Path.home() / ".local/share"))
    folder = root / "ActsReports"
    folder.mkdir(parents=True, exist_ok=True)
    return folder / "imports.sqlite3"


class ImportDatabase:
    def __init__(self, path: Path | None = None) -> None:
        self.path = path or database_path()
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self.connect() as db:
            db.execute(
                """
                CREATE TABLE IF NOT EXISTS imports (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    source_name TEXT NOT NULL,
                    source_path TEXT NOT NULL,
                    sha256 TEXT NOT NULL UNIQUE,
                    imported_at TEXT NOT NULL,
                    workbook BLOB NOT NULL,
                    network TEXT NOT NULL,
                    period TEXT NOT NULL
                )
                """
            )

    def connect(self) -> sqlite3.Connection:
        db = sqlite3.connect(self.path)
        db.row_factory = sqlite3.Row
        return db

    def add_workbook(
        self,
        source: Path,
        content: bytes,
        network: str,
        period: str,
    ) -> tuple[int, bool]:
        digest = hashlib.sha256(content).hexdigest()
        with self.connect() as db:
            previous = db.execute(
                "SELECT id FROM imports WHERE sha256 = ?", (digest,)
            ).fetchone()
            if previous:
                return int(previous["id"]), False
            cursor = db.execute(
                """
                INSERT INTO imports
                    (source_name, source_path, sha256, imported_at, workbook, network, period)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    source.name,
                    str(source.resolve()),
                    digest,
                    datetime.now().isoformat(timespec="seconds"),
                    sqlite3.Binary(content),
                    network,
                    period,
                ),
            )
            return int(cursor.lastrowid), True

    def get_workbook(self, import_id: int) -> bytes:
        with self.connect() as db:
            row = db.execute(
                "SELECT workbook FROM imports WHERE id = ?", (import_id,)
            ).fetchone()
        if row is None:
            raise ValueError("Не знайдено імпортовану книгу в локальній базі.")
        return bytes(row["workbook"])


@dataclass
class WorkbookInfo:
    networks: list[str]
    network: str
    period: str
    sheet_names: set[str]


def _read_cell_value(cell: object) -> object:
    return getattr(cell, "value", None)


def _display_period(value: object, fallback: object = None) -> str:
    if isinstance(value, str) and value.strip():
        return value.strip()
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, (int, float)):
        try:
            from openpyxl.utils.datetime import from_excel

            return from_excel(value).strftime("%d.%m.%Y")
        except (OverflowError, ValueError):
            pass
    if isinstance(fallback, str) and fallback.strip():
        return fallback.strip()
    return "Період не визначено"


def inspect_workbook(content: bytes) -> WorkbookInfo:
    try:
        workbook = _load_workbook_values(content)
    except (OSError, ValueError, zipfile.BadZipFile, InvalidFileException) as exc:
        raise ValueError(f"Не вдалося відкрити Excel-книгу: {exc}") from exc
    try:
        names = set(workbook.sheetnames)
        missing = sorted(REQUIRED_SHEETS - names)
        if missing:
            raise ValueError(
                "У книзі відсутні обов’язкові аркуші: " + ", ".join(missing)
            )

        pivot = workbook["PIVOT"]
        selected_network = _read_cell_value(pivot["B5"])
        selected_network = str(selected_network).strip() if selected_network else ""
        if not selected_network:
            raise ValueError("Не знайдено мережу у клітинці PIVOT!B5.")

        period = _display_period(
            _read_cell_value(pivot["F1"]), _read_cell_value(pivot["E1"])
        )

        has_selected_network = False
        for sheet_name in ("Реквізити", "Реквізити_лік"):
            sheet = workbook[sheet_name]
            for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2, values_only=True):
                value = row[0]
                if value is not None and str(value).strip() == selected_network:
                    has_selected_network = True
        if not has_selected_network:
            raise ValueError(
                f"Мережа «{selected_network}» з PIVOT!B5 відсутня у довідниках реквізитів."
            )

        return WorkbookInfo(
            networks=[selected_network],
            network=selected_network,
            period=period,
            sheet_names=names,
        )
    finally:
        workbook.close()


def _copy_sheet_values(
    source: object,
    destination: object,
    *,
    title: str,
) -> None:
    destination.title = title
    destination.sheet_view.showGridLines = False
    destination.freeze_panes = "A2"
    destination.sheet_properties.pageSetUpPr.fitToPage = True
    destination.page_setup.fitToWidth = 1
    destination.page_setup.fitToHeight = 0
    destination.page_setup.orientation = "landscape"
    destination.page_margins.left = 0.25
    destination.page_margins.right = 0.25
    destination.page_margins.top = 0.5
    destination.page_margins.bottom = 0.5

    for row_index, row in enumerate(source.iter_rows(), start=1):
        has_value = False
        for col_index, source_cell in enumerate(row, start=1):
            value = source_cell.value
            if value is None:
                continue
            has_value = True
            cell = destination.cell(row_index, col_index, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_index <= 4:
                cell.font = Font(name="Arial", size=10, bold=True)
                cell.fill = SUBHEADER_FILL
            if row_index in (9, 10) and title.startswith("Звіт"):
                cell.font = WHITE_BOLD
                cell.fill = HEADER_FILL
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            if isinstance(value, (int, float)):
                cell.number_format = '#,##0.00;[Red]-#,##0.00;–'
            cell.border = Border(bottom=THIN_GREY)
        if has_value:
            destination.row_dimensions[row_index].height = 24 if row_index <= 10 else 20

    if destination.max_column:
        for column in range(1, destination.max_column + 1):
            letter = get_column_letter(column)
            if title == "Акт":
                width = 100 if column == 1 else 20 if column < 4 else 15
            else:
                width = 8 if column == 1 else 48 if column == 2 else 18 if column == 3 else 14
            destination.column_dimensions[letter].width = width
    if title.startswith("Звіт"):
        destination.auto_filter.ref = destination.dimensions
    destination.sheet_properties.pageSetUpPr.fitToPage = True


def create_preview_workbook(
    content: bytes,
    variant: str,
    network: str,
    period: str,
) -> tuple[Workbook, list[str]]:
    if variant not in VARIANTS:
        raise ValueError("Не вибрано варіант ЛІК або ДД.")
    try:
        source = _load_workbook_values(content)
    except (OSError, ValueError, zipfile.BadZipFile, InvalidFileException) as exc:
        raise ValueError(f"Не вдалося прочитати імпортовану книгу: {exc}") from exc

    selected = VARIANTS[variant]
    required = (selected["act"], selected["report"], selected["requisites"])
    missing = [name for name in required if name not in source.sheetnames]
    if missing:
        source.close()
        raise ValueError("Відсутні аркуші для вибраного варіанта: " + ", ".join(missing))

    actual_network = source["PIVOT"]["B5"].value
    if str(actual_network).strip() != network:
        source.close()
        raise ValueError(
            "Дані PIVOT підготовлені для мережі "
            f"«{actual_network}», а вибрано «{network}». "
            "Оновіть/відфільтруйте PIVOT у вихідній книзі та імпортуйте її повторно."
        )

    warnings: list[str] = []
    requisites = source[selected["requisites"]]
    match_rows = [
        row
        for row in requisites.iter_rows(min_row=2, values_only=True)
        if len(row) > 1 and str(row[1]).strip() == network
    ]
    if not match_rows:
        source.close()
        raise ValueError(
            f"У «{selected['requisites']}» немає реквізитів для мережі «{network}»."
        )
    if len(match_rows) > 1:
        warnings.append(
            f"У «{selected['requisites']}» знайдено {len(match_rows)} рядки мережі; "
            "у вихідній книзі формули використовують перший збіг."
        )

    output = Workbook()
    act = output.active
    report = output.create_sheet()
    try:
        _copy_sheet_values(source[selected["act"]], act, title="Акт")
        _copy_sheet_values(source[selected["report"]], report, title="Звіт з викладки")
        act_total = act["C14"].value
        act_base = act["C12"].value
        act_extra = act["C13"].value
        report_total = report[selected["report_total"]].value
        if not all(
            isinstance(value, (int, float))
            for value in (act_total, act_base, act_extra, report_total)
        ):
            raise ValueError(
                "У книзі немає збережених числових підсумків акта або звіту. "
                "Відкрийте її в Excel, оновіть розрахунки та імпортуйте повторно."
            )
        if (
            abs(Decimal(str(act_base)) - Decimal(str(report_total))) > Decimal("0.01")
            or abs(
                Decimal(str(act_total))
                - Decimal(str(act_base))
                - Decimal(str(act_extra))
            )
            > Decimal("0.01")
        ):
            raise ValueError(
                "Підсумок звіту не збігається з актом. Перевірте перерахунок "
                "формул в Excel і дані зведених таблиць перед імпортом."
            )
        act.sheet_properties.tabColor = "1F4E78"
        report.sheet_properties.tabColor = "70AD47"
        output.properties.title = f"Акт і звіт з викладки — {network}"
        output.properties.subject = f"{variant}; {period}"
        output.properties.creator = APP_NAME
    finally:
        source.close()
    return output, warnings


class ActsReportsApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title(APP_NAME)
        self.root.geometry("900x620")
        self.root.minsize(760, 500)
        self.database = ImportDatabase()
        self.import_id: int | None = None
        self.info: WorkbookInfo | None = None
        self.source_path: Path | None = None
        self.preview_import_id: int | None = None
        self.preview_network = ""
        self.preview_variant = ""
        self._build_ui()

    def _build_ui(self) -> None:
        outer = ttk.Frame(self.root, padding=16)
        outer.pack(fill="both", expand=True)

        ttk.Label(outer, text=APP_NAME, font=("Segoe UI", 18, "bold")).pack(
            anchor="w"
        )
        ttk.Label(
            outer,
            text="Завантажте книгу, виберіть мережу та тип документа, "
            "перегляньте аркуші й збережіть результат.",
            wraplength=840,
        ).pack(anchor="w", pady=(4, 14))

        file_row = ttk.Frame(outer)
        file_row.pack(fill="x")
        self.file_label = ttk.Label(file_row, text="Файл не вибрано", anchor="w")
        self.file_label.pack(side="left", fill="x", expand=True)
        ttk.Button(file_row, text="Завантажити Excel…", command=self.choose_file).pack(
            side="right"
        )

        options = ttk.LabelFrame(outer, text="Параметри документа", padding=10)
        options.pack(fill="x", pady=12)
        ttk.Label(options, text="Мережа:").grid(row=0, column=0, sticky="w")
        self.network_var = tk.StringVar()
        self.network_combo = ttk.Combobox(
            options, textvariable=self.network_var, state="readonly", width=42
        )
        self.network_combo.grid(row=0, column=1, sticky="w", padx=(8, 18))
        self.network_combo.bind("<<ComboboxSelected>>", self._network_selected)
        ttk.Label(options, text="Варіант:").grid(row=0, column=2, sticky="w")
        self.variant_var = tk.StringVar(value="ЛІК")
        self.variant_combo = ttk.Combobox(
            options,
            textvariable=self.variant_var,
            state="readonly",
            values=list(VARIANTS),
            width=10,
        )
        self.variant_combo.grid(row=0, column=3, sticky="w", padx=8)
        self.variant_combo.bind("<<ComboboxSelected>>", self._options_changed)
        ttk.Label(options, text="Період:").grid(
            row=1, column=0, sticky="w", pady=(10, 0)
        )
        self.period_label = ttk.Label(options, text="—")
        self.period_label.grid(
            row=1, column=1, columnspan=3, sticky="w", padx=(8, 0), pady=(10, 0)
        )
        ttk.Label(
            options,
            text="За замовчуванням використовується весь період, наявний у PIVOT.",
        ).grid(row=2, column=0, columnspan=4, sticky="w", pady=(8, 0))

        action_row = ttk.Frame(outer)
        action_row.pack(fill="x", pady=(0, 8))
        self.preview_button = ttk.Button(
            action_row,
            text="Попередній перегляд",
            command=self.preview,
            state="disabled",
        )
        self.preview_button.pack(side="left")
        self.save_button = ttk.Button(
            action_row, text="Зберегти Excel…", command=self.save_result, state="disabled"
        )
        self.save_button.pack(side="left", padx=8)
        self.status_var = tk.StringVar(value="Очікується завантаження Excel-файлу.")
        ttk.Label(action_row, textvariable=self.status_var).pack(
            side="left", padx=10, fill="x", expand=True
        )

        preview_box = ttk.LabelFrame(outer, text="Попередній перегляд", padding=8)
        preview_box.pack(fill="both", expand=True)
        self.preview_tabs = ttk.Notebook(preview_box)
        self.preview_tabs.pack(fill="both", expand=True)
        self.preview_tabs.bind("<<NotebookTabChanged>>", self._tab_changed)
        self.preview_tree: dict[str, ttk.Treeview] = {}
        self.preview_tabs.add(ttk.Frame(self.preview_tabs), text="Акт")
        self.preview_tabs.add(ttk.Frame(self.preview_tabs), text="Звіт з викладки")

    def choose_file(self) -> None:
        selected = filedialog.askopenfilename(
            title="Оберіть вхідну книгу",
            filetypes=[("Книги Excel", "*.xlsx *.xlsm"), ("Усі файли", "*.*")],
        )
        if not selected:
            return
        path = Path(selected)
        try:
            content = path.read_bytes()
            info = inspect_workbook(content)
            import_id, inserted = self.database.add_workbook(
                path, content, info.network, info.period
            )
        except (OSError, sqlite3.Error, ValueError) as exc:
            messagebox.showerror("Не вдалося завантажити файл", str(exc), parent=self.root)
            self.status_var.set("Помилка імпорту.")
            return

        self.import_id = import_id
        self.info = info
        self.source_path = path
        self.file_label.configure(text=path.name)
        self.network_combo.configure(values=info.networks)
        self.network_var.set(info.network)
        self.period_label.configure(text=info.period)
        self.preview_button.configure(state="normal")
        self.save_button.configure(state="disabled")
        self._clear_preview()
        self.preview_import_id = None
        self.status_var.set(
            "Книгу збережено в локальній базі."
            if inserted
            else "Ця книга вже була завантажена; використовується збережений імпорт."
        )

    def _network_selected(self, _event: object = None) -> None:
        self._options_changed()

    def _options_changed(self, _event: object = None) -> None:
        self.save_button.configure(state="disabled")
        self.preview_import_id = None
        self._clear_preview()
        if self.info and self.network_var.get() != self.info.network:
            self.status_var.set(
                "Попередній перегляд буде доступний лише для мережі, "
                "на яку відфільтрований PIVOT у цій книзі."
            )
        elif self.info:
            self.status_var.set("Мережу вибрано.")

    def _clear_preview(self) -> None:
        for widget in self.preview_tabs.winfo_children():
            self.preview_tabs.forget(widget)
            widget.destroy()
        self.preview_tree.clear()
        self.preview_tabs.add(ttk.Frame(self.preview_tabs), text="Акт")
        self.preview_tabs.add(ttk.Frame(self.preview_tabs), text="Звіт з викладки")

    def _populate_tree(self, tab_index: int, sheet: object, title: str) -> None:
        frame = self.preview_tabs.winfo_children()[tab_index]
        wrapper = ttk.Frame(frame)
        wrapper.pack(fill="both", expand=True)
        columns = tuple(f"c{index}" for index in range(1, sheet.max_column + 1))
        tree = ttk.Treeview(wrapper, columns=columns, show="headings")
        yscroll = ttk.Scrollbar(wrapper, orient="vertical", command=tree.yview)
        xscroll = ttk.Scrollbar(wrapper, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        tree.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        xscroll.grid(row=1, column=0, sticky="ew")
        wrapper.rowconfigure(0, weight=1)
        wrapper.columnconfigure(0, weight=1)
        for index, column in enumerate(columns, start=1):
            tree.heading(column, text=get_column_letter(index))
            tree.column(column, width=180 if index == 2 else 120, stretch=True)
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=1, max_col=len(columns), values_only=True), start=1
        ):
            if not any(value is not None for value in row):
                continue
            rendered = tuple(
                str(value).replace("\n", " ")[:240] if value is not None else ""
                for value in row
            )
            tree.insert("", "end", values=rendered, text=str(row_number))
        self.preview_tree[title] = tree

    def preview(self) -> None:
        if self.import_id is None:
            return
        try:
            content = self.database.get_workbook(self.import_id)
            workbook, warnings = create_preview_workbook(
                content,
                self.variant_var.get(),
                self.network_var.get(),
                self.info.period if self.info else "",
            )
            self._clear_preview()
            self._populate_tree(0, workbook["Акт"], "Акт")
            self._populate_tree(1, workbook["Звіт з викладки"], "Звіт з викладки")
            self.preview_tabs.select(0)
            self.preview_import_id = self.import_id
            self.preview_network = self.network_var.get()
            self.preview_variant = self.variant_var.get()
            self.save_button.configure(state="normal")
            self.status_var.set("Перегляд готовий. Дані документа відповідають імпорту.")
            if warnings:
                messagebox.showwarning(
                    "Перевірте реквізити", "\n".join(warnings), parent=self.root
                )
            workbook.close()
        except (OSError, sqlite3.Error, ValueError, KeyError) as exc:
            self.save_button.configure(state="disabled")
            self.status_var.set("Перевірка не пройдена.")
            messagebox.showerror("Не вдалося підготувати перегляд", str(exc), parent=self.root)

    def _tab_changed(self, _event: object = None) -> None:
        selected = self.preview_tabs.select()
        if selected:
            self.root.title(f"{APP_NAME} — {self.preview_tabs.tab(selected, 'text')}")

    def save_result(self) -> None:
        if self.import_id is None or self.info is None:
            return
        if (
            self.preview_import_id != self.import_id
            or self.preview_network != self.network_var.get()
            or self.preview_variant != self.variant_var.get()
        ):
            self.save_button.configure(state="disabled")
            messagebox.showerror(
                "Спершу перегляньте документи",
                "Перед збереженням сформуйте попередній перегляд для вибраних параметрів.",
                parent=self.root,
            )
            return
        default_name = f"Акт_звіт_{self.network_var.get()}_{self.variant_var.get()}.xlsx"
        destination = filedialog.asksaveasfilename(
            title="Зберегти акт і звіт",
            defaultextension=".xlsx",
            initialfile=default_name,
            filetypes=[("Книга Excel", "*.xlsx")],
        )
        if not destination:
            return
        try:
            content = self.database.get_workbook(self.import_id)
            workbook, warnings = create_preview_workbook(
                content,
                self.variant_var.get(),
                self.network_var.get(),
                self.info.period,
            )
            workbook.save(destination)
            workbook.close()
        except (OSError, sqlite3.Error, ValueError, KeyError) as exc:
            messagebox.showerror("Не вдалося зберегти Excel", str(exc), parent=self.root)
            return
        self.status_var.set(f"Файл збережено: {destination}")
        message = f"Створено книгу з двома аркушами:\n{destination}"
        if warnings:
            message += "\n\nПопередження:\n" + "\n".join(warnings)
        messagebox.showinfo("Готово", message, parent=self.root)


def main() -> None:
    root = tk.Tk()
    try:
        ttk.Style(root).theme_use("vista" if os.name == "nt" else "clam")
    except tk.TclError:
        pass
    ActsReportsApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
